from docx import Document

doc = Document()

doc.add_heading("RAG Engineering Notes", level=1)

doc.add_paragraph(
    "Retrieval-Augmented Generation (RAG) retrieves relevant "
    "information from external knowledge sources to help an LLM "
    "answer questions accurately."
)

doc.add_heading("Key Decisions", level=2)
doc.add_paragraph(
    "Support multiple document formats.",
    style="List Bullet",
)
doc.add_paragraph(
    "Validate extracted content before indexing.",
    style="List Bullet",
)
doc.add_paragraph(
    "Record ingestion errors without stopping the entire batch.",
    style="List Bullet",
)

doc.save("datasets/sample/meeting_notes.docx")
print("Created a valid DOCX sample.")